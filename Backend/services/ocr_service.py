import os
import csv

from pypdf import PdfReader


def extract_text(file_path: str) -> str:

    if not os.path.exists(file_path):

        return "File not found."


    extension = os.path.splitext(
        file_path
    )[1].lower()


    # -----------------------------------------
    # TXT
    # -----------------------------------------

    if extension == ".txt":

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8"
            ) as file:

                return file.read()

        except Exception as error:

            return (
                f"Text extraction error: {str(error)}"
            )


    # -----------------------------------------
    # CSV
    # -----------------------------------------

    if extension == ".csv":

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8-sig",
                newline=""
            ) as file:

                return "\n".join(
                    "\t".join(row)
                    for row in csv.reader(file)
                )

        except Exception as error:

            return f"CSV extraction error: {str(error)}"


    # -----------------------------------------
    # Excel workbooks
    # -----------------------------------------

    if extension == ".xlsx":

        try:

            from openpyxl import load_workbook

            workbook = load_workbook(
                file_path,
                read_only=True,
                data_only=True
            )

            return "\n".join(
                "\t".join("" if value is None else str(value) for value in row)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
            )

        except Exception as error:

            return f"XLSX extraction error: {str(error)}"

    if extension == ".xls":

        try:

            import xlrd

            workbook = xlrd.open_workbook(file_path, on_demand=True)

            return "\n".join(
                "\t".join(str(value) for value in sheet.row_values(row_index))
                for sheet in workbook.sheets()
                for row_index in range(sheet.nrows)
            )

        except Exception as error:

            return f"XLS extraction error: {str(error)}"


    # -----------------------------------------
    # PDF
    # -----------------------------------------

    if extension == ".pdf":

        try:

            reader = PdfReader(
                file_path
            )

            pages_text = []


            for page in reader.pages:

                text = page.extract_text()

                if text:

                    pages_text.append(text)


            extracted_text = "\n".join(
                pages_text
            )


            if extracted_text.strip():

                return extracted_text


            return (
                "No digital text found in PDF. "
                "OCR is required for this scanned PDF."
            )


        except Exception as error:

            return (
                f"PDF extraction error: {str(error)}"
            )


    # -----------------------------------------
    # Images
    # -----------------------------------------

    if extension in {
        ".png",
        ".jpg",
        ".jpeg"
    }:

        return (
            "Image uploaded successfully. "
            "OCR processing for images will be "
            "connected next."
        )


    # -----------------------------------------
    # DOC/DOCX
    # -----------------------------------------

    if extension in {
        ".doc",
        ".docx"
    }:

        return (
            "Document uploaded successfully. "
            "DOC/DOCX text extraction will be "
            "connected next."
        )


    return (
        "File uploaded, but text extraction "
        "is not available for this file type."
    )